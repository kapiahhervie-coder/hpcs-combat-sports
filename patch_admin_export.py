"""
patch_admin_export.py
Menambahkan fitur Export ke Excel ke combat/admin.py secara otomatis:
  1. Update import (tambah SpeedAgilityAuditL4 + library Excel)
  2. Sisipkan fungsi export_ke_excel SEBELUM class AtletAdmin
  3. Tambahkan actions = [export_ke_excel] di dalam AtletAdmin
Aman dijalankan berkali-kali (idempotent) -- melewati bagian yang sudah benar.
"""
import os
import shutil
from datetime import datetime

FILEPATH = os.path.join('combat', 'admin.py')

OLD_IMPORT = """from django.contrib import admin
from django.utils.html import format_html
from .models import (
    Atlet,
    CorrectionAuditL1,
    StrengthAuditL2,
    PowerAuditL3,
    RekomendasiProgram,
)"""

NEW_IMPORT = """from django.contrib import admin
from django.utils.html import format_html
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill
from .models import (
    Atlet,
    CorrectionAuditL1,
    StrengthAuditL2,
    PowerAuditL3,
    SpeedAgilityAuditL4,
    RekomendasiProgram,
)"""


EXPORT_FUNCTION = '''

# ══════════════════════════════════════════════════════════════════════
# EXPORT KE EXCEL — Admin Action (hanya bisa diakses via Django Admin)
# ══════════════════════════════════════════════════════════════════════

def export_ke_excel(modeladmin, request, queryset):
    """
    Export semua data audit (L1 boxing generic, L1 Muay Thai khusus,
    L2, L3, L4) untuk atlet yang dipilih -- jadi 1 file Excel,
    1 sheet, 1 baris per record audit.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HPCS Data Atlet"

    headers = ['Nama Atlet', 'Cabang', 'Kategori', 'Gender', 'Level', 'Tanggal', 'Predikat', 'Total Skor', 'Detail Skor', 'Catatan']
    ws.append(headers)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    try:
        from muaythai.models import CorrectionAuditL1MT
    except ImportError:
        CorrectionAuditL1MT = None

    for atlet in queryset:
        for a in CorrectionAuditL1.objects.filter(atlet=atlet):
            detail = (
                f"Rotation:{a.score_rotation} | Extension:{a.score_extension} | "
                f"Stability:{a.score_stability} | Posture:{a.score_posture} | "
                f"Breathing:{a.score_breathing}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L1 - Correction', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

        if CorrectionAuditL1MT:
            for a in CorrectionAuditL1MT.objects.filter(atlet=atlet):
                detail = (
                    f"HipRotation:{a.skor_neural_hip_rotasi} | Balance:{a.skor_neural_balance} | "
                    f"Ankle:{a.skor_neural_ankle} | Thoracic:{a.skor_neural_thoracic} | "
                    f"Shoulder:{a.skor_neural_shoulder} | HipHinge:{a.skor_neural_hip_hinge} | "
                    f"Core:{a.skor_neural_core} | Recovery:{a.skor_neural_recovery}"
                )
                ws.append([
                    atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                    'L1 - Mobility/Stability (MT)', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                    a.predikat, a.total_skor, detail, a.catatan or '',
                ])

        for a in StrengthAuditL2.objects.filter(atlet=atlet):
            detail = (
                f"Lower:{a.score_lower} | Push:{a.score_push} | Pull:{a.score_pull} | "
                f"Core:{a.score_core} | Isometric:{a.score_isometric} | "
                f"TremorRatio:{a.iso_tremor_rasio}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L2 - Strength', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

        for a in PowerAuditL3.objects.filter(atlet=atlet):
            detail = (
                f"Jump:{a.score_jump} | Sprint:{a.score_sprint} | Throw:{a.score_throw} | "
                f"RSI:{a.score_rsi} | Agility:{a.score_agility}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L3 - Power', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

        for a in SpeedAgilityAuditL4.objects.filter(atlet=atlet):
            detail = (
                f"Hex:{a.score_hex} | Punch:{a.score_punch} | Yoyo:{a.score_yoyo} | "
                f"KondisiUji:{a.kondisi_uji}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L4 - Speed & Agility', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

    for col_cells in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 60)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="HPCS_Data_Atlet.xlsx"'
    wb.save(response)
    return response

export_ke_excel.short_description = "\U0001F4CA Export atlet terpilih ke Excel"


class AtletAdmin(admin.ModelAdmin):'''

OLD_CLASS_MARKER = "class AtletAdmin(admin.ModelAdmin):"


OLD_ORDERING = "    ordering      = ['nama_atlet']\n    # Hanya data atlet"
NEW_ORDERING = "    ordering      = ['nama_atlet']\n    actions       = [export_ke_excel]\n    # Hanya data atlet"


def main():
    if not os.path.exists(FILEPATH):
        print(f"File tidak ditemukan: {FILEPATH}")
        return

    with open(FILEPATH, encoding='utf-8') as f:
        content = f.read()

    original_content = content
    applied = []
    skipped = []
    warned = []

    # 1. Import
    if content.count(OLD_IMPORT) == 1:
        content = content.replace(OLD_IMPORT, NEW_IMPORT)
        applied.append("Update import")
    elif content.count(NEW_IMPORT) >= 1:
        skipped.append("Update import (sudah benar)")
    else:
        warned.append("Update import (blok lama tidak ditemukan persis)")

    # 2. Sisipkan fungsi export SEBELUM class AtletAdmin
    if "def export_ke_excel" in content:
        skipped.append("Fungsi export_ke_excel (sudah ada)")
    elif content.count(OLD_CLASS_MARKER) == 1:
        content = content.replace(OLD_CLASS_MARKER, EXPORT_FUNCTION, 1)
        applied.append("Sisipkan fungsi export_ke_excel")
    else:
        warned.append(f"Sisipkan fungsi export (marker class AtletAdmin ditemukan {content.count(OLD_CLASS_MARKER)}x, perlu 1x)")

    # 3. Tambahkan actions di AtletAdmin
    if "actions       = [export_ke_excel]" in content:
        skipped.append("actions = [export_ke_excel] (sudah ada)")
    elif content.count(OLD_ORDERING) == 1:
        content = content.replace(OLD_ORDERING, NEW_ORDERING)
        applied.append("Tambah actions = [export_ke_excel]")
    else:
        warned.append("Tambah actions (marker ordering+comment tidak ditemukan persis)")

    if content == original_content:
        print("Tidak ada perubahan dilakukan.")
    else:
        timestamp = datetime.now().strftime('%Y%m%d-%H%M%S')
        backup_path = FILEPATH + f'.backup-{timestamp}'
        shutil.copy2(FILEPATH, backup_path)
        print(f"Backup dibuat: {backup_path}\n")

        with open(FILEPATH, 'w', encoding='utf-8', newline='') as f:
            f.write(content)
        print(f"File ditulis ulang: {FILEPATH}\n")

    print("=== RINGKASAN ===")
    for a in applied:
        print(f"  [OK] {a}")
    for s in skipped:
        print(f"  [-]  {s}")
    for w in warned:
        print(f"  [!]  {w}")

    print("\nJalankan 'python manage.py check' untuk verifikasi.")


if __name__ == '__main__':
    main()
