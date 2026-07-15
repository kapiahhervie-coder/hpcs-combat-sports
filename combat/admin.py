"""
HPCS Combat Sports - Admin
Filosofi: Data atlet input SEKALI, terbawa otomatis ke semua level audit.
"""
from django.contrib import admin
from django.utils.html import format_html
from .models import (
    Atlet,
    CorrectionAuditL1,
    StrengthAuditL2,
    PowerAuditL3,
    SpeedAgilityAuditL4,
    RekomendasiProgram,
)
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill


# ── Helper badge ──────────────────────────────────────────────────────

def badge_predikat(predikat):
    warna = {'ELITE':'#34d399','READY':'#60a5fa','DEVELOPING':'#fbbf24','NOVICE':'#f87171'}
    bg    = {'ELITE':'rgba(52,211,153,0.15)','READY':'rgba(96,165,250,0.15)',
             'DEVELOPING':'rgba(251,191,36,0.15)','NOVICE':'rgba(248,113,113,0.15)'}
    c = warna.get(predikat, '#888')
    b = bg.get(predikat, 'transparent')
    return format_html(
        '<span style="background:{};color:{};border:1px solid {};'
        'border-radius:20px;padding:2px 10px;font-weight:700;font-size:0.8rem;">{}</span>',
        b, c, c, predikat
    )


def badge_layak(value):
    from django.utils.safestring import mark_safe
    if value:
        return mark_safe('<span style="color:#34d399;font-weight:700;">&#10003; LAYAK</span>')
    return mark_safe('<span style="color:#f87171;font-weight:700;">&#10007; BELUM</span>')

# ══════════════════════════════════════════════════════════════════════
# ATLET — master data, input sekali
# ══════════════════════════════════════════════════════════════════════



# ══════════════════════════════════════════════════════════════════════
# EXPORT KE EXCEL — Admin Action (hanya bisa diakses via Django Admin)
# ══════════════════════════════════════════════════════════════════════

def export_ke_excel(modeladmin, request, queryset):
    """
    Export semua data audit (L1 boxing generic, L1 Muay Thai khusus,
    L2, L3, L4) untuk atlet yang dipilih -- jadi 1 file Excel,
    1 sheet, 1 baris per record audit.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HPCS Data Atlet"

    headers = ['Nama Atlet', 'Cabang', 'Kategori', 'Gender', 'Level', 'Tanggal', 'Predikat', 'Total Skor', 'Detail Skor', 'Catatan']
    ws.append(headers)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    try:
        from muaythai.models import CorrectionAuditL1MT
    except ImportError:
        CorrectionAuditL1MT = None

    for atlet in queryset:
        for a in CorrectionAuditL1.objects.filter(atlet=atlet):
            detail = (
                f"Rotation:{a.score_rotation} | Extension:{a.score_extension} | "
                f"Stability:{a.score_stability} | Posture:{a.score_posture} | "
                f"Breathing:{a.score_breathing}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L1 - Correction', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

        if CorrectionAuditL1MT:
            for a in CorrectionAuditL1MT.objects.filter(atlet=atlet):
                detail = (
                    f"HipRotation:{a.skor_neural_hip_rotasi} | Balance:{a.skor_neural_balance} | "
                    f"Ankle:{a.skor_neural_ankle} | Thoracic:{a.skor_neural_thoracic} | "
                    f"Shoulder:{a.skor_neural_shoulder} | HipHinge:{a.skor_neural_hip_hinge} | "
                    f"Core:{a.skor_neural_core} | Recovery:{a.skor_neural_recovery}"
                )
                ws.append([
                    atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                    'L1 - Mobility/Stability (MT)', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                    a.predikat, a.total_skor, detail, a.catatan or '',
                ])

        for a in StrengthAuditL2.objects.filter(atlet=atlet):
            detail = (
                f"Lower:{a.score_lower} | Push:{a.score_push} | Pull:{a.score_pull} | "
                f"Core:{a.score_core} | Isometric:{a.score_isometric} | "
                f"TremorRatio:{a.iso_tremor_rasio}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L2 - Strength', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

        for a in PowerAuditL3.objects.filter(atlet=atlet):
            detail = (
                f"Jump:{a.score_jump} | Sprint:{a.score_sprint} | Throw:{a.score_throw} | "
                f"RSI:{a.score_rsi} | Agility:{a.score_agility}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L3 - Power', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

        for a in SpeedAgilityAuditL4.objects.filter(atlet=atlet):
            detail = (
                f"Hex:{a.score_hex} | Punch:{a.score_punch} | Yoyo:{a.score_yoyo} | "
                f"KondisiUji:{a.kondisi_uji}"
            )
            ws.append([
                atlet.nama_atlet, atlet.cabang, a.kategori_usia, a.gender,
                'L4 - Speed & Agility', a.timestamp.strftime('%d/%m/%Y %H:%M'),
                a.predikat, a.total_skor, detail, a.catatan or '',
            ])

    for col_cells in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 3, 60)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="HPCS_Data_Atlet.xlsx"'
    wb.save(response)
    return response

export_ke_excel.short_description = "📊 Export atlet terpilih ke Excel"


@admin.register(Atlet)
class AtletAdmin(admin.ModelAdmin):
    list_display  = ['nama_atlet', 'gender', 'kategori_umur', 'kelas_berat', 'cabang']
    list_filter   = ['gender', 'kategori_umur', 'cabang']
    search_fields = ['nama_atlet']
    ordering      = ['nama_atlet']
    actions       = [export_ke_excel]

    # Hanya data atlet — tidak ada field audit di sini
    fields = ['nama_atlet', 'gender', 'kategori_umur', 'kelas_berat', 'tanggal_lahir', 'cabang']


# ══════════════════════════════════════════════════════════════════════
# AUDIT L1 — pilih atlet + input skor saja
# ══════════════════════════════════════════════════════════════════════

@admin.register(CorrectionAuditL1)
class CorrectionAuditL1Admin(admin.ModelAdmin):
    list_display  = [
        'get_atlet_nama', 'get_atlet_gender', 'get_atlet_bb',
        'score_rotation', 'score_extension', 'score_stability',
        'score_posture', 'score_breathing',
        'get_total', 'get_predikat', 'get_layak', 'timestamp',
    ]
    list_filter   = ['predikat', 'layak_naik']
    search_fields = ['atlet__nama_atlet']
    ordering      = ['-timestamp']
    readonly_fields = ['total_skor', 'predikat', 'layak_naik', 'alasan_tidak_layak']

    # Form input: pilih atlet + 5 skor + catatan
    # Data atlet (gender, BB, dll) otomatis terbawa dari FK
    fields = [
        'atlet',                          # ← dropdown pilih atlet
        'score_rotation', 'score_extension', 'score_stability',
        'score_posture', 'score_breathing',
        'catatan',
        # Hasil otomatis — tidak perlu diisi manual
        'total_skor', 'predikat', 'layak_naik', 'alasan_tidak_layak',
    ]

    def get_atlet_nama(self, obj):
        return obj.atlet.nama_atlet if obj.atlet else obj.atlet_name
    get_atlet_nama.short_description = 'Nama Atlet'

    def get_atlet_gender(self, obj):
        return obj.atlet.gender if obj.atlet else '—'
    get_atlet_gender.short_description = 'Gender'

    def get_atlet_bb(self, obj):
        return f"{obj.atlet.kelas_berat} kg" if obj.atlet else '—'
    get_atlet_bb.short_description = 'BB'

    def get_total(self, obj):
        c = '#34d399' if obj.total_skor >= 7 else ('#fbbf24' if obj.total_skor >= 5 else '#f87171')
        return format_html('<strong style="color:{};">{}</strong>', c, obj.total_skor)
    get_total.short_description = 'Total'

    def get_predikat(self, obj):
        return badge_predikat(obj.predikat)
    get_predikat.short_description = 'Predikat'

    def get_layak(self, obj):
        return badge_layak(obj.layak_naik)
    get_layak.short_description = 'Layak ke L2?'

    def save_model(self, request, obj, form, change):
        # Otomatis isi data dari atlet saat save lewat admin
        if obj.atlet:
            obj.atlet_name  = obj.atlet.nama_atlet
            obj.gender      = obj.atlet.gender
            obj.kelas_berat = obj.atlet.kelas_berat
            obj.kategori_usia = obj.atlet.kategori_umur
        super().save_model(request, obj, form, change)


# ══════════════════════════════════════════════════════════════════════
# AUDIT L2 — pilih atlet + input skor + data mentah
# ══════════════════════════════════════════════════════════════════════

@admin.register(StrengthAuditL2)
class StrengthAuditL2Admin(admin.ModelAdmin):
    list_display  = [
        'get_atlet_nama', 'get_atlet_gender', 'get_atlet_bb',
        'score_lower', 'score_push', 'score_pull',
        'score_core', 'score_isometric',
        'get_tremor', 'get_total', 'get_predikat', 'get_layak', 'timestamp',
    ]
    list_filter   = ['predikat', 'layak_naik']
    search_fields = ['atlet__nama_atlet', 'atlet_name']
    ordering      = ['-timestamp']
    readonly_fields = [
        'lower_5rm_bw_ratio', 'iso_tremor_rasio',
        'total_skor', 'predikat', 'layak_naik', 'alasan_tidak_layak',
    ]

    fields = [
        'atlet',                          # ← dropdown pilih atlet
        # Pilar 1
        'lower_5rm_beban', 'lower_5rm_bw_ratio', 'score_lower',
        # Pilar 2
        'push_5rm_beban', 'score_push',
        # Pilar 3
        'pull_reps', 'score_pull',
        # Pilar 4
        'core_durasi_detik', 'score_core',
        # Pilar 5 + CNS
        'iso_durasi_detik', 'iso_tremor_onset_detik', 'iso_tremor_rasio', 'score_isometric',
        'catatan',
        # Hasil otomatis
        'total_skor', 'predikat', 'layak_naik', 'alasan_tidak_layak',
    ]

    def get_atlet_nama(self, obj):
        return obj.atlet.nama_atlet if obj.atlet else obj.atlet_name
    get_atlet_nama.short_description = 'Nama Atlet'

    def get_atlet_gender(self, obj):
        return obj.atlet.gender if obj.atlet else '—'
    get_atlet_gender.short_description = 'Gender'

    def get_atlet_bb(self, obj):
        return f"{obj.atlet.kelas_berat} kg" if obj.atlet else '—'
    get_atlet_bb.short_description = 'BB'

    def get_tremor(self, obj):
        if obj.iso_tremor_rasio is None:
            return '—'
        c = '#34d399' if obj.iso_tremor_rasio >= 85 else ('#fbbf24' if obj.iso_tremor_rasio >= 65 else '#f87171')
        return format_html('<span style="color:{};font-weight:700;">{}%</span>', c, obj.iso_tremor_rasio)
    get_tremor.short_description = 'CNS Tremor'

    def get_total(self, obj):
        c = '#34d399' if obj.total_skor >= 7 else ('#fbbf24' if obj.total_skor >= 5 else '#f87171')
        return format_html('<strong style="color:{};">{}</strong>', c, obj.total_skor)
    get_total.short_description = 'Total'

    def get_predikat(self, obj):
        return badge_predikat(obj.predikat)
    get_predikat.short_description = 'Predikat'

    def get_layak(self, obj):
        return badge_layak(obj.layak_naik)
    get_layak.short_description = 'Layak ke L3?'

    def save_model(self, request, obj, form, change):
        # Otomatis isi data dari atlet saat save lewat admin
        if obj.atlet:
            obj.atlet_name    = obj.atlet.nama_atlet
            obj.gender        = obj.atlet.gender
            obj.kelas_berat   = obj.atlet.kelas_berat
            obj.kategori_usia = obj.atlet.kategori_umur
        super().save_model(request, obj, form, change)


# ══════════════════════════════════════════════════════════════════════
# AUDIT L3 — pilih atlet + input skor
# ══════════════════════════════════════════════════════════════════════

@admin.register(PowerAuditL3)
class PowerAuditL3Admin(admin.ModelAdmin):
    list_display  = [
        'get_atlet_nama', 'get_atlet_gender', 'get_atlet_bb',
        'score_jump', 'score_sprint', 'score_throw',
        'score_rsi', 'score_agility',
        'get_total', 'get_predikat', 'get_layak', 'timestamp',
    ]
    list_filter   = ['predikat', 'layak_naik']
    search_fields = ['atlet__nama_atlet', 'atlet_name']
    ordering      = ['-timestamp']
    readonly_fields = [
        'rsi_value', 'total_skor', 'predikat',
        'layak_naik', 'layak_bertahan', 'alasan_tidak_layak',
    ]

    fields = [
        'atlet',                          # ← dropdown pilih atlet
        'score_jump', 'score_sprint', 'score_throw',
        'rsi_jump_height', 'rsi_contact_time', 'rsi_value', 'score_rsi',
        'score_agility',
        'catatan',
        # Hasil otomatis
        'total_skor', 'predikat', 'layak_naik', 'layak_bertahan', 'alasan_tidak_layak',
    ]

    def get_atlet_nama(self, obj):
        return obj.atlet.nama_atlet if obj.atlet else obj.atlet_name
    get_atlet_nama.short_description = 'Nama Atlet'

    def get_atlet_gender(self, obj):
        return obj.atlet.gender if obj.atlet else '—'
    get_atlet_gender.short_description = 'Gender'

    def get_atlet_bb(self, obj):
        return f"{obj.atlet.kelas_berat} kg" if obj.atlet else '—'
    get_atlet_bb.short_description = 'BB'

    def get_total(self, obj):
        c = '#34d399' if obj.total_skor >= 9 else ('#60a5fa' if obj.total_skor >= 7 else ('#fbbf24' if obj.total_skor >= 5 else '#f87171'))
        return format_html('<strong style="color:{};">{}</strong>', c, obj.total_skor)
    get_total.short_description = 'Total'

    def get_predikat(self, obj):
        return badge_predikat(obj.predikat)
    get_predikat.short_description = 'Predikat'

    def get_layak(self, obj):
        return badge_layak(obj.layak_naik)
    get_layak.short_description = 'Layak Kompetisi?'

    def save_model(self, request, obj, form, change):
        # Otomatis isi data dari atlet saat save lewat admin
        if obj.atlet:
            obj.atlet_name    = obj.atlet.nama_atlet
            obj.gender        = obj.atlet.gender
            obj.kelas_berat   = obj.atlet.kelas_berat
            obj.kategori_usia = obj.atlet.kategori_umur
        super().save_model(request, obj, form, change)


@admin.register(SpeedAgilityAuditL4)
class SpeedAgilityAuditL4Admin(admin.ModelAdmin):
    list_display  = [
        'get_atlet_nama', 'get_atlet_gender', 'get_atlet_bb',
        'score_hex', 'score_punch', 'score_yoyo',
        'get_total', 'get_predikat', 'get_layak', 'timestamp',
    ]
    list_filter   = ['predikat', 'layak_kompetisi', 'kondisi_uji']
    search_fields = ['atlet__nama_atlet', 'atlet_name']
    ordering      = ['-timestamp']
    readonly_fields = [
        'hex_waktu_rata', 'yoyo_vo2max_estimasi', 'total_skor', 'predikat',
        'layak_kompetisi', 'layak_bertahan', 'rekomendasi_auto', 'alasan_tidak_layak',
    ]

    fields = [
        'atlet',
        'hex_waktu_putaran1', 'hex_waktu_putaran2', 'hex_waktu_putaran3',
        'hex_waktu_rata', 'score_hex',
        'punch_freq_10s', 'punch_postur_ok', 'punch_reaction_time_ms', 'score_punch',
        'yoyo_level_tercapai', 'yoyo_shuttle_tercapai', 'yoyo_total_jarak_m',
        'yoyo_vo2max_estimasi', 'score_yoyo',
        'kondisi_uji', 'catatan',
        'total_skor', 'predikat', 'layak_kompetisi', 'layak_bertahan',
        'rekomendasi_auto', 'alasan_tidak_layak',
    ]

    def get_atlet_nama(self, obj):
        return obj.atlet.nama_atlet if obj.atlet else obj.atlet_name
    get_atlet_nama.short_description = 'Nama Atlet'

    def get_atlet_gender(self, obj):
        return obj.atlet.gender if obj.atlet else '\u2014'
    get_atlet_gender.short_description = 'Gender'

    def get_atlet_bb(self, obj):
        return f"{obj.atlet.kelas_berat} kg" if obj.atlet else '\u2014'
    get_atlet_bb.short_description = 'BB'

    def get_total(self, obj):
        c = '#34d399' if obj.total_skor >= 9 else ('#60a5fa' if obj.total_skor >= 7 else ('#fbbf24' if obj.total_skor >= 5 else '#f87171'))
        return format_html('<strong style="color:{};">{}</strong>', c, obj.total_skor)
    get_total.short_description = 'Total'

    def get_predikat(self, obj):
        return badge_predikat(obj.predikat)
    get_predikat.short_description = 'Predikat'

    def get_layak(self, obj):
        return badge_layak(obj.layak_kompetisi)
    get_layak.short_description = 'Layak Kompetisi?'

    def save_model(self, request, obj, form, change):
        if obj.atlet:
            obj.atlet_name    = obj.atlet.nama_atlet
            obj.gender        = obj.atlet.gender
            obj.kelas_berat   = obj.atlet.kelas_berat
            obj.kategori_usia = obj.atlet.kategori_umur
        super().save_model(request, obj, form, change)


# ══════════════════════════════════════════════════════════════════════
# REKOMENDASI PROGRAM
# ══════════════════════════════════════════════════════════════════════

@admin.register(RekomendasiProgram)
class RekomendasiProgramAdmin(admin.ModelAdmin):
    list_display  = ['atlet', 'audit_level', 'tanggal', 'durasi_minggu', 'sudah_dijalankan']
    list_filter   = ['audit_level', 'sudah_dijalankan']
    search_fields = ['atlet__nama_atlet']
    fields        = ['atlet', 'audit_level', 'tanggal', 'isi_rekomendasi', 'durasi_minggu', 'sudah_dijalankan']
