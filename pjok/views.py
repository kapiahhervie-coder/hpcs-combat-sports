import json
from datetime import date
from django.http import HttpResponse
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from django.contrib.auth.decorators import login_required
from django.db import models
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from .models import Siswa, FASE_CHOICES, JENJANG_PER_FASE, SesiAbsensi, Absensi, RencanaMingguan

from .diagnostik import (
    KOMPONEN_LABEL,
    diagnosa_akar_masalah,
    rekomendasi_perbaikan,
    rubrik_label,
)
from .forms import EditSiswaForm, GuruProfileForm, PenilaianFisikForm, PenilaianKarakterForm, PenilaianPengetahuanForm, PenilaianTeknikForm, RencanaMingguanForm, SiswaForm
from .models import GuruProfile, MateriFase, Siswa


@login_required
def dashboard_redirect(request):
    """Titik masuk setelah login — arahkan ke dashboard fase guru, atau ke setup profil."""
    try:
        profile = request.user.guruprofile
    except GuruProfile.DoesNotExist:
        return redirect('pjok:buat_profile')
    return redirect('pjok:dashboard_fase', fase=profile.fase)


@login_required
def buat_profile(request):
    """Setup profil guru — pilih fase yang diampu."""
    if request.method == 'POST':
        form = GuruProfileForm(request.POST)
        if form.is_valid():
            profile = form.save(commit=False)
            profile.user = request.user
            profile.save()
            return redirect('pjok:dashboard_fase', fase=profile.fase)
    else:
        form = GuruProfileForm()
    return render(request, 'pjok/buat_profile.html', {'form': form})


@login_required
def dashboard_fase(request, fase):
    """
    Dashboard utama guru untuk fase tertentu.
    Sementara: semua guru bisa melihat semua fase (akan dibatasi kembali nanti).
    """
    profile = request.user.guruprofile

    siswa_list = Siswa.objects.filter(guru=profile, fase=fase)

    # Hitung skor kesiapan terbaru per siswa untuk badge di tabel & statistik ringkas
    siswa_data = []
    skor_semua = []
    tier_count = {'SANGAT BAIK': 0, 'BAIK': 0, 'CUKUP': 0, 'PERLU PERHATIAN': 0}

    for s in siswa_list:
        fisik = s.penilaian_fisik.order_by('-tanggal_tes').first()
        kesiapan = None
        tier = None
        if fisik and fisik.skor_l1 is not None:
            kesiapan = round((fisik.skor_l1 + fisik.skor_l2 + fisik.skor_l3 + fisik.skor_l4) / 4, 1)
            skor_semua.append(kesiapan)
            if kesiapan >= 85:
                tier = 'SANGAT BAIK'
            elif kesiapan >= 70:
                tier = 'BAIK'
            elif kesiapan >= 55:
                tier = 'CUKUP'
            else:
                tier = 'PERLU PERHATIAN'
            tier_count[tier] += 1
        siswa_data.append({'siswa': s, 'kesiapan': kesiapan, 'tier': tier})

    rata_rata_kesiapan = round(sum(skor_semua) / len(skor_semua), 1) if skor_semua else None
    total_baik_ke_atas = tier_count['SANGAT BAIK'] + tier_count['BAIK']
    total_perlu_perhatian = tier_count['PERLU PERHATIAN']

    selected_siswa_id = request.GET.get('siswa_id')
    selected_siswa = None
    if selected_siswa_id:
        selected_siswa = siswa_list.filter(id=selected_siswa_id).first()

    form_fisik = PenilaianFisikForm()
    form_teknik = PenilaianTeknikForm(fase=fase)
    form_karakter = PenilaianKarakterForm()
    form_pengetahuan = PenilaianPengetahuanForm()

    return render(request, 'pjok/dashboard.html', {
        'siswa_data': siswa_data,
        'total_siswa': siswa_list.count(),
        'rata_rata_kesiapan': rata_rata_kesiapan,
        'total_baik_ke_atas': total_baik_ke_atas,
        'total_perlu_perhatian': total_perlu_perhatian,
        'fase': fase,
        'jenjang': JENJANG_PER_FASE.get(fase, ''),
        'profile': profile,
        'fase_list': [
            {'kode': k, 'label': v, 'jenjang': JENJANG_PER_FASE.get(k, ''), 'aktif': k == fase}
            for k, v in FASE_CHOICES
        ],
        'siswa_list': siswa_list,
        'selected_siswa': selected_siswa,
        'form_fisik': form_fisik,
        'form_teknik': form_teknik,
        'form_karakter': form_karakter,
        'form_pengetahuan': form_pengetahuan,
    })


@login_required
def tambah_siswa(request):
    """Form tambah siswa baru — terikat ke guru yang login, fase diambil dari halaman asal."""
    profile = request.user.guruprofile
    fase_target = request.GET.get('fase') or profile.fase

    if request.method == 'POST':
        form = SiswaForm(request.POST)
        if form.is_valid():
            siswa = form.save(commit=False)
            siswa.guru = profile
            siswa.fase = request.POST.get('fase') or fase_target
            siswa.save()
            return redirect('pjok:dashboard_fase', fase=siswa.fase)
    else:
        form = SiswaForm()

    return render(request, 'pjok/tambah_siswa.html', {
        'form': form,
        'fase': fase_target,
        'jenjang': JENJANG_PER_FASE.get(fase_target, ''),
    })


@login_required
def rencana_mingguan(request, fase):
    """Kelola rencana materi mingguan untuk Prota/Prosem, per tahun ajaran & semester."""
    profile = request.user.guruprofile
    tahun_ajaran = request.GET.get('tahun', '2026/2027')
    semester = request.GET.get('semester', 'ganjil')

    if request.method == 'POST':
        minggu_ke = request.POST.get('minggu_ke')
        instance = RencanaMingguan.objects.filter(
            guru=profile, fase=fase, tahun_ajaran=tahun_ajaran,
            semester=semester, minggu_ke=minggu_ke,
        ).first()
        form = RencanaMingguanForm(request.POST, instance=instance, fase=fase)
        if form.is_valid():
            rencana = form.save(commit=False)
            rencana.guru = profile
            rencana.fase = fase
            rencana.tahun_ajaran = tahun_ajaran
            rencana.semester = semester
            rencana.save()
        return redirect(f"{request.path}?tahun={tahun_ajaran}&semester={semester}")

    rencana_list = RencanaMingguan.objects.filter(
        guru=profile, fase=fase, tahun_ajaran=tahun_ajaran, semester=semester,
    ).select_related('materi')
    rencana_map = {r.minggu_ke: r for r in rencana_list}

    JUMLAH_MINGGU = 18
    rows = []
    for i in range(1, JUMLAH_MINGGU + 1):
        rows.append({'minggu_ke': i, 'rencana': rencana_map.get(i)})

    return render(request, 'pjok/rencana_mingguan.html', {
        'rows': rows,
        'fase': fase,
        'jenjang': JENJANG_PER_FASE.get(fase, ''),
        'tahun_ajaran': tahun_ajaran,
        'semester': semester,
        'form_materi_qs': MateriFase.objects.filter(fase=fase),
    })


@login_required
def daftar_absensi(request, fase):
    """Daftar semua sesi absensi yang pernah diambil untuk fase ini."""
    profile = request.user.guruprofile
    sesi_list = SesiAbsensi.objects.filter(guru=profile, fase=fase).annotate(
        total_hadir=models.Count('daftar_absensi', filter=models.Q(daftar_absensi__status='H')),
        total_siswa=models.Count('daftar_absensi'),
    )
    return render(request, 'pjok/daftar_absensi.html', {
        'sesi_list': sesi_list,
        'fase': fase,
        'jenjang': JENJANG_PER_FASE.get(fase, ''),
    })


@login_required
def rekap_absensi(request, fase):
    """Tabel rekap: baris siswa, kolom tanggal sesi, isi status kehadiran."""
    profile = request.user.guruprofile
    siswa_list = Siswa.objects.filter(guru=profile, fase=fase).order_by('nama')
    sesi_list = SesiAbsensi.objects.filter(guru=profile, fase=fase).order_by('tanggal')

    absen_qs = Absensi.objects.filter(sesi__in=sesi_list).select_related('sesi', 'siswa')
    absen_map = {}
    for a in absen_qs:
        absen_map[(a.siswa_id, a.sesi_id)] = a

    rows = []
    for s in siswa_list:
        cells = []
        rekap = {'H': 0, 'I': 0, 'S': 0, 'A': 0}
        for sesi in sesi_list:
            a = absen_map.get((s.id, sesi.id))
            status = a.status if a else None
            if status:
                rekap[status] += 1
            cells.append({'sesi': sesi, 'status': status})
        rows.append({'siswa': s, 'cells': cells, 'rekap': rekap})

    return render(request, 'pjok/rekap_absensi.html', {
        'rows': rows,
        'sesi_list': sesi_list,
        'fase': fase,
        'jenjang': JENJANG_PER_FASE.get(fase, ''),
    })


@login_required
def ambil_absensi(request, fase):
    """Ambil/edit absensi untuk fase & tanggal tertentu (?tanggal=YYYY-MM-DD, default hari ini)."""
    profile = request.user.guruprofile
    tanggal_str = request.GET.get('tanggal') or date.today().isoformat()
    tanggal = date.fromisoformat(tanggal_str)

    siswa_list = Siswa.objects.filter(guru=profile, fase=fase).order_by('nama')
    sesi, _ = SesiAbsensi.objects.get_or_create(guru=profile, fase=fase, tanggal=tanggal)

    if request.method == 'POST':
        sesi.catatan_sesi = request.POST.get('catatan_sesi', '')
        sesi.save()
        for s in siswa_list:
            status = request.POST.get(f'status_{s.id}', 'H')
            keterangan = request.POST.get(f'keterangan_{s.id}', '')
            Absensi.objects.update_or_create(
                sesi=sesi, siswa=s,
                defaults={'status': status, 'keterangan': keterangan},
            )
        return redirect('pjok:daftar_absensi', fase=fase)

    absen_map = {a.siswa_id: a for a in sesi.daftar_absensi.all()}
    rows = [{'siswa': s, 'absen': absen_map.get(s.id)} for s in siswa_list]

    return render(request, 'pjok/ambil_absensi.html', {
        'rows': rows,
        'sesi': sesi,
        'tanggal': tanggal,
        'fase': fase,
        'jenjang': JENJANG_PER_FASE.get(fase, ''),
        'status_choices': Absensi.STATUS_CHOICES,
    })


@login_required
def edit_siswa(request, siswa_id):
    """Guru bisa mengedit data siswa miliknya sendiri, termasuk memindahkan fase."""
    profile = request.user.guruprofile
    siswa = get_object_or_404(Siswa, id=siswa_id, guru=profile)

    if request.method == 'POST':
        form = EditSiswaForm(request.POST, instance=siswa)
        if form.is_valid():
            siswa = form.save()
            return redirect('pjok:dashboard_fase', fase=siswa.fase)
    else:
        form = EditSiswaForm(instance=siswa)

    return render(request, 'pjok/edit_siswa.html', {'form': form, 'siswa': siswa})


@login_required
def tambah_penilaian_fisik(request, siswa_id):
    """Form input tes fisik mentah — skor L1-L4 dihitung otomatis lewat save() di models.py."""
    profile = request.user.guruprofile
    siswa = get_object_or_404(Siswa, id=siswa_id, guru=profile)  # pastikan siswa milik guru ini

    if request.method == 'POST':
        form = PenilaianFisikForm(request.POST)
        if form.is_valid():
            penilaian = form.save(commit=False)
            penilaian.siswa = siswa
            penilaian.save()
            return redirect('pjok:detail_siswa', siswa_id=siswa.id)
    else:
        form = PenilaianFisikForm()

    return render(request, 'pjok/tambah_penilaian_fisik.html', {'form': form, 'siswa': siswa})


@login_required
def tambah_penilaian_teknik(request, siswa_id):
    """Form input nilai teknik/skill per materi — dipakai untuk diagnosis akar masalah."""
    profile = request.user.guruprofile
    siswa = get_object_or_404(Siswa, id=siswa_id, guru=profile)  # pastikan siswa milik guru ini

    if request.method == 'POST':
        form = PenilaianTeknikForm(request.POST, fase=profile.fase)
        if form.is_valid():
            penilaian = form.save(commit=False)
            penilaian.siswa = siswa
            penilaian.save()
            return redirect('pjok:detail_siswa', siswa_id=siswa.id)
    else:
        form = PenilaianTeknikForm(fase=profile.fase)

    return render(request, 'pjok/tambah_penilaian_teknik.html', {'form': form, 'siswa': siswa})


@login_required
def tambah_penilaian_karakter(request, siswa_id):
    """Form input nilai karakter per aspek (Sportivitas, Kerja Sama, dll)."""
    profile = request.user.guruprofile
    siswa = get_object_or_404(Siswa, id=siswa_id, guru=profile)

    if request.method == 'POST':
        form = PenilaianKarakterForm(request.POST)
        if form.is_valid():
            penilaian = form.save(commit=False)
            penilaian.siswa = siswa
            penilaian.save()
            return redirect('pjok:detail_siswa', siswa_id=siswa.id)
    else:
        form = PenilaianKarakterForm()

    return render(request, 'pjok/tambah_penilaian_karakter.html', {'form': form, 'siswa': siswa})


@login_required
def tambah_penilaian_pengetahuan(request, siswa_id):
    """Form input nilai pengetahuan gerak (Elemen 2 CP PJOK)."""
    profile = request.user.guruprofile
    siswa = get_object_or_404(Siswa, id=siswa_id, guru=profile)

    if request.method == 'POST':
        form = PenilaianPengetahuanForm(request.POST)
        if form.is_valid():
            penilaian = form.save(commit=False)
            penilaian.siswa = siswa
            penilaian.save()
            return redirect('pjok:detail_siswa', siswa_id=siswa.id)
    else:
        form = PenilaianPengetahuanForm()

    return render(request, 'pjok/tambah_penilaian_pengetahuan.html', {'form': form, 'siswa': siswa})


@login_required
def detail_siswa(request, siswa_id):
    """
    Laporan diagnostik siswa — mengikuti pola 'Athlete Performance Report' di combat
    sports (radar chart, trend chart, kartu kekuatan/keterbatasan/rekomendasi),
    diadaptasi untuk konteks PJOK: L1-L4 fisik, nilai teknik, diagnosis akar masalah.
    """
    profile = request.user.guruprofile
    siswa = get_object_or_404(Siswa, id=siswa_id, guru=profile)  # pastikan siswa milik guru ini

    AMBANG = 70
    riwayat_fisik = list(siswa.penilaian_fisik.order_by('tanggal_tes'))
    fisik_terbaru = riwayat_fisik[-1] if riwayat_fisik else None
    teknik_list = siswa.penilaian_teknik.select_related('materi').order_by('-tanggal')
    karakter_list = siswa.penilaian_karakter.order_by('-tanggal')
    pengetahuan_list = siswa.penilaian_pengetahuan.order_by('-tanggal')

    radar_data_json = None
    trend_data_json = None
    months_labels_json = json.dumps([])
    rubrik = {}
    rekomendasi = []
    kekuatan = []
    keterbatasan = []
    kesiapan = None

    if fisik_terbaru:
        skor_map = {
            'l1': fisik_terbaru.skor_l1 or 0,
            'l2': fisik_terbaru.skor_l2 or 0,
            'l3': fisik_terbaru.skor_l3 or 0,
            'l4': fisik_terbaru.skor_l4 or 0,
        }
        kesiapan = round(sum(skor_map.values()) / 4, 1)

        radar_data_json = json.dumps({
            'labels': ['Core & Mobility', 'Strength', 'Power', 'Speed & Agility'],
            'data': [skor_map['l1'], skor_map['l2'], skor_map['l3'], skor_map['l4']],
        })
        rubrik = {
            'l1': rubrik_label(fisik_terbaru.skor_l1),
            'l2': rubrik_label(fisik_terbaru.skor_l2),
            'l3': rubrik_label(fisik_terbaru.skor_l3),
            'l4': rubrik_label(fisik_terbaru.skor_l4),
        }
        rekomendasi = rekomendasi_perbaikan(fisik_terbaru)

        for kode, label in KOMPONEN_LABEL.items():
            if skor_map[kode] >= AMBANG:
                kekuatan.append({'label': label, 'skor': skor_map[kode]})
            else:
                keterbatasan.append({'label': label, 'skor': skor_map[kode]})

        # Tren historis dari seluruh riwayat tes fisik siswa ini
        months_labels_json = json.dumps([r.tanggal_tes.strftime('%d %b') for r in riwayat_fisik])
        trend_data_json = json.dumps({
            'l1': [r.skor_l1 or 0 for r in riwayat_fisik],
            'l2': [r.skor_l2 or 0 for r in riwayat_fisik],
            'l3': [r.skor_l3 or 0 for r in riwayat_fisik],
            'l4': [r.skor_l4 or 0 for r in riwayat_fisik],
        })

    diagnosis_per_materi = []
    materi_dinilai = MateriFase.objects.filter(
        penilaian_teknik__siswa=siswa
    ).distinct()
    for materi in materi_dinilai:
        hasil = diagnosa_akar_masalah(siswa, materi, ambang_batas=AMBANG)
        if hasil:
            diagnosis_per_materi.append({'materi': materi, 'diagnosis': hasil})

    if kesiapan is None:
        tier = None
    elif kesiapan >= 85:
        tier = 'SANGAT BAIK'
    elif kesiapan >= 70:
        tier = 'BAIK'
    elif kesiapan >= 55:
        tier = 'CUKUP'
    else:
        tier = 'PERLU PERHATIAN'

    umur = None
    if siswa.tanggal_lahir:
        hari_ini = date.today()
        umur = hari_ini.year - siswa.tanggal_lahir.year - (
            (hari_ini.month, hari_ini.day) < (siswa.tanggal_lahir.month, siswa.tanggal_lahir.day)
        )

    profil_dominan = None
    if fisik_terbaru:
        semua = kekuatan + keterbatasan
        if semua:
            profil_dominan = max(semua, key=lambda x: x['skor'])['label']

    return render(request, 'pjok/detail_siswa.html', {
        'siswa': siswa,
        'profile': profile,
        'umur': umur,
        'fisik_terbaru': fisik_terbaru,
        'jumlah_tes': len(riwayat_fisik),
        'teknik_list': teknik_list,
        'karakter_list': karakter_list,
        'pengetahuan_list': pengetahuan_list,
        'radar_data_json': radar_data_json,
        'trend_data_json': trend_data_json,
        'months_labels_json': months_labels_json,
        'rubrik': rubrik,
        'rekomendasi': rekomendasi,
        'kekuatan': kekuatan,
        'keterbatasan': keterbatasan,
        'kesiapan': kesiapan,
        'tier': tier,
        'profil_dominan': profil_dominan,
        'diagnosis_per_materi': diagnosis_per_materi,
        'tanggal_cetak': timezone.now(),
    })


def _bikin_dokumen_dasar(judul, guru, fase, jenjang):
    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    h = doc.add_heading(judul, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Mata Pelajaran: PJOK (Pendidikan Jasmani, Olahraga, dan Kesehatan)")
    run.font.size = Pt(11)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(f"Fase {fase} ({jenjang}) &mdash; Sekolah: {guru.sekolah}").font.size = Pt(11)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run(f"Guru Pengampu: {guru.nama_lengkap}").font.size = Pt(11)

    doc.add_paragraph()
    return doc


def _isi_tabel_rencana(doc, rows, dengan_semester=False):
    kolom = 5 if dengan_semester else 4
    table = doc.add_table(rows=1, cols=kolom)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = 'Minggu'
    hdr[1].text = 'Materi Pembelajaran'
    hdr[2].text = 'Alokasi JP'
    hdr[3].text = 'Keterangan'
    if dengan_semester:
        hdr[4].text = 'Semester'

    for row in rows:
        cells = table.add_row().cells
        cells[0].text = str(row['minggu_ke'])
        if row.get('rencana'):
            cells[1].text = row['rencana'].nama_tampil
            cells[2].text = str(row['rencana'].alokasi_jp)
            cells[3].text = row['rencana'].keterangan or ''
        else:
            cells[1].text = '-'
            cells[2].text = '-'
            cells[3].text = ''
        if dengan_semester:
            cells[4].text = row.get('semester_label', '')


@login_required
def export_prosem(request, fase):
    """Generate dokumen Word Program Semester (Prosem)."""
    profile = request.user.guruprofile
    tahun_ajaran = request.GET.get('tahun', '2026/2027')
    semester = request.GET.get('semester', 'ganjil')

    rencana_list = RencanaMingguan.objects.filter(
        guru=profile, fase=fase, tahun_ajaran=tahun_ajaran, semester=semester,
    ).select_related('materi')
    rencana_map = {r.minggu_ke: r for r in rencana_list}
    rows = [{'minggu_ke': i, 'rencana': rencana_map.get(i)} for i in range(1, 19)]

    judul_semester = 'Ganjil' if semester == 'ganjil' else 'Genap'
    doc = _bikin_dokumen_dasar(
        f"PROGRAM SEMESTER {judul_semester.upper()}\nTahun Ajaran {tahun_ajaran}",
        profile, fase, JENJANG_PER_FASE.get(fase, ''),
    )
    _isi_tabel_rencana(doc, rows)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="Prosem_{judul_semester}_{fase}_{tahun_ajaran.replace("/", "-")}.docx"'
    doc.save(response)
    return response


@login_required
def export_prota(request, fase):
    """Generate dokumen Word Program Tahunan (Prota), gabungan semester ganjil + genap."""
    profile = request.user.guruprofile
    tahun_ajaran = request.GET.get('tahun', '2026/2027')

    doc = _bikin_dokumen_dasar(
        f"PROGRAM TAHUNAN\nTahun Ajaran {tahun_ajaran}",
        profile, fase, JENJANG_PER_FASE.get(fase, ''),
    )

    for semester, label in [('ganjil', 'Ganjil'), ('genap', 'Genap')]:
        doc.add_heading(f"Semester {label}", level=2)
        rencana_list = RencanaMingguan.objects.filter(
            guru=profile, fase=fase, tahun_ajaran=tahun_ajaran, semester=semester,
        ).select_related('materi')
        rencana_map = {r.minggu_ke: r for r in rencana_list}
        rows = [{'minggu_ke': i, 'rencana': rencana_map.get(i)} for i in range(1, 19)]
        _isi_tabel_rencana(doc, rows)
        doc.add_paragraph()

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="Prota_{fase}_{tahun_ajaran.replace("/", "-")}.docx"'
    doc.save(response)
    return response